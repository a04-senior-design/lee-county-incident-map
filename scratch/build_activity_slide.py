from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

NAVY   = RGBColor(0x1E, 0x3A, 0x5F)
GREY   = RGBColor(0x44, 0x44, 0x44)
BLUE_FILL   = RGBColor(0xDC, 0xE6, 0xF1)
GREEN_FILL  = RGBColor(0xE2, 0xEF, 0xDA)
YELLOW_FILL = RGBColor(0xFF, 0xF2, 0xCC)
RED_FILL    = RGBColor(0xF8, 0xCE, 0xCC)
BLACK       = RGBColor(0x20, 0x20, 0x20)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])


def add_shape(shape_type, left, top, width, height, text="", fill=WHITE,
              font_size=10, bold=False, font_color=BLACK, line_color=RGBColor(0x60, 0x60, 0x60)):
    shp = slide.shapes.add_shape(shape_type, Inches(left), Inches(top), Inches(width), Inches(height))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line_color
    shp.line.width = Pt(1)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(45720)
    tf.margin_right = Emu(45720)
    tf.margin_top = Emu(18000)
    tf.margin_bottom = Emu(18000)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(font_size)
        p.font.bold = bold
        p.font.color.rgb = font_color
    return shp


def center(shp):
    return (shp.left / 914400 + shp.width / 914400 / 2,
            shp.top / 914400 + shp.height / 914400 / 2)


def edge_point(shp, side):
    l, t = shp.left / 914400, shp.top / 914400
    w, h = shp.width / 914400, shp.height / 914400
    if side == "top":    return (l + w / 2, t)
    if side == "bottom": return (l + w / 2, t + h)
    if side == "left":   return (l, t + h / 2)
    if side == "right":  return (l + w, t + h / 2)


def arrow(p1, p2, color=RGBColor(0x40, 0x40, 0x40)):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                       Inches(p1[0]), Inches(p1[1]),
                                       Inches(p2[0]), Inches(p2[1]))
    conn.line.color.rgb = color
    conn.line.width = Pt(1.25)
    ln = conn.line._get_or_add_ln()
    tail = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
    ln.append(tail)
    return conn


def label(text, x, y, size=8, color=GREY, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(0.5), Inches(0.22))
    tf = box.text_frame
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    tf.text = text
    p = tf.paragraphs[0]
    p.font.size = Pt(size)
    p.font.italic = True
    p.font.bold = bold
    p.font.color.rgb = color


# ── Title ─────────────────────────────────────────────────────────────────
title_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.25), Inches(12.5), Inches(0.6))
tf = title_box.text_frame
tf.text = "Generate Report — UI Activity Flow"
p = tf.paragraphs[0]
p.font.size = Pt(26); p.font.bold = True; p.font.color.rgb = NAVY

sub_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.82), Inches(12.5), Inches(0.4))
tf2 = sub_box.text_frame
tf2.text = "Design approach and proposed implementation for the Frontend / Backend request lifecycle"
p2 = tf2.paragraphs[0]
p2.font.size = Pt(14); p2.font.italic = True; p2.font.color.rgb = GREY

# ── Left column: bullets ─────────────────────────────────────────────────
bullets_box = slide.shapes.add_textbox(Inches(0.4), Inches(1.35), Inches(4.65), Inches(5.9))
btf = bullets_box.text_frame
btf.word_wrap = True

sections = [
    ("Design Approach", [
        "Single synchronous request/response — the whole flow (read bounds "
        "→ POST → filter → render → assemble → download) "
        "completes in one HTTP round trip; no polling or websockets.",
        "Button-state guard — the Generate Report button disables and relabels "
        "to “Generating…” immediately on click to block duplicate submissions.",
        "Fail-fast on the server — zero incidents surviving the bounds+date "
        "filters returns HTTP 400 before any map or chart rendering is attempted.",
        "Symmetric restore path — success and failure both converge on one "
        "“restore button” step via a finally block, so the UI can never "
        "get stuck loading.",
    ]),
    ("Proposed Implementation", [
        "Frontend reads map.getBounds() + the active date filter, POSTs "
        "{days, bounds} JSON to /api/report.",
        "Backend pulls the in-memory geocoded incident cache and applies "
        "sequential viewport, then date-window filters.",
        "On a match: _map_snapshot() (staticmap) and _bar_chart()/_line_chart() "
        "(Matplotlib) render into BytesIO buffers, then generate_pdf() "
        "(ReportLab) assembles and streams the PDF.",
        "On no match: HTTP 400 returns immediately, skipping tile-download "
        "and chart-rendering cost entirely.",
        "Browser receives the PDF blob, creates a temporary object URL, "
        "triggers download via a synthetic anchor click, and revokes the URL.",
    ]),
]

first = True
for heading, items in sections:
    hp = btf.paragraphs[0] if first else btf.add_paragraph()
    first = False
    hp.text = heading
    hp.font.size = Pt(15); hp.font.bold = True; hp.font.color.rgb = NAVY
    hp.space_after = Pt(4)
    for item in items:
        ip = btf.add_paragraph()
        ip.text = "•  " + item
        ip.font.size = Pt(10.5); ip.font.color.rgb = GREY
        ip.space_after = Pt(6)
        ip.level = 1
    if heading != sections[-1][0]:
        sp = btf.add_paragraph()
        sp.space_after = Pt(6)

# ── Right column: activity diagram ──────────────────────────────────────────
FE_L, FE_W = 5.3, 3.6
BE_L, BE_W = 9.1, 3.6
FE_C = FE_L + FE_W / 2
BE_C = BE_L + BE_W / 2

add_shape(MSO_SHAPE.RECTANGLE, FE_L, 1.30, FE_W, 0.35, "Frontend (Browser)", BLUE_FILL, 11, True, NAVY)
add_shape(MSO_SHAPE.RECTANGLE, BE_L, 1.30, BE_W, 0.35, "Backend (Flask Server)", GREEN_FILL, 11, True, NAVY)

divider = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(9.0), Inches(1.30), Inches(9.0), Inches(6.35))
divider.line.color.rgb = RGBColor(0x99, 0x99, 0x99)
divider.line.width = Pt(1)

fe1 = add_shape(MSO_SHAPE.OVAL, FE_C - 0.15, 1.80, 0.3, 0.3, "", BLACK)
fe2 = add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 5.6, 2.25, 3.0, 0.5, "User clicks\n“Generate Report”", BLUE_FILL)
fe3 = add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 5.6, 2.90, 3.0, 0.65, "Read map.getBounds() &\ndate filter; POST /api/report", BLUE_FILL)
fe4 = add_shape(MSO_SHAPE.DIAMOND, 5.9, 3.70, 2.4, 0.55, "Response\nOK?", YELLOW_FILL, 9)
fe_no = add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 5.6, 4.40, 1.4, 0.5, "Display\nerror alert", RED_FILL, 9)
fe_yes = add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 7.15, 4.40, 1.75, 0.65, "Receive PDF blob,\ntrigger download", BLUE_FILL, 9)
fe_merge = add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 5.8, 5.20, 2.6, 0.5, "Restore button", BLUE_FILL)
fe_end = add_shape(MSO_SHAPE.OVAL, FE_C - 0.175, 5.85, 0.35, 0.35, "", BLACK)

be1 = add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 9.4, 1.75, 3.0, 0.5, "Parse JSON body\n(days, bounds)", GREEN_FILL)
be2 = add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 9.4, 2.40, 3.0, 0.6, "Retrieve incidents;\napply viewport + date filters", GREEN_FILL, 9)
be3 = add_shape(MSO_SHAPE.DIAMOND, 9.6, 3.15, 2.6, 0.6, "Incidents\nremaining?", YELLOW_FILL, 9)
be_no = add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 9.3, 3.90, 1.3, 0.45, "Return\nHTTP 400", RED_FILL, 9)
be_yes = add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 10.75, 3.90, 1.85, 0.65, "Render snapshot &\ngenerate charts", GREEN_FILL, 9)
be5 = add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 9.5, 4.70, 2.6, 0.5, "Assemble PDF\n(ReportLab)", GREEN_FILL)
be6 = add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 9.5, 5.35, 2.6, 0.6, "Return\napplication/pdf response", GREEN_FILL, 9)

# ── Frontend-lane arrows ─────────────────────────────────────────────────
arrow(edge_point(fe1, "bottom"), edge_point(fe2, "top"))
arrow(edge_point(fe2, "bottom"), edge_point(fe3, "top"))
arrow(edge_point(fe3, "bottom"), edge_point(fe4, "top"))
arrow(edge_point(fe4, "left"), edge_point(fe_no, "top"))
arrow(edge_point(fe4, "right"), edge_point(fe_yes, "top"))
arrow(edge_point(fe_no, "bottom"), edge_point(fe_merge, "top"))
arrow(edge_point(fe_yes, "bottom"), edge_point(fe_merge, "top"))
arrow(edge_point(fe_merge, "bottom"), edge_point(fe_end, "top"))
label("No", 5.75, 4.10, bold=True)
label("Yes", 8.10, 4.10, bold=True)

# ── Backend-lane arrows ──────────────────────────────────────────────────
arrow(edge_point(be1, "bottom"), edge_point(be2, "top"))
arrow(edge_point(be2, "bottom"), edge_point(be3, "top"))
arrow(edge_point(be3, "left"), edge_point(be_no, "top"))
arrow(edge_point(be3, "right"), edge_point(be_yes, "top"))
arrow(edge_point(be_yes, "bottom"), edge_point(be5, "top"))
arrow(edge_point(be5, "bottom"), edge_point(be6, "top"))
label("No (None)", 9.15, 3.62, bold=True)
label("Yes (Some)", 11.55, 3.62, bold=True)

# ── Cross-lane arrows ─────────────────────────────────────────────────────
arrow(edge_point(fe3, "right"), edge_point(be1, "left"), color=NAVY)
arrow(edge_point(be_no, "bottom"), edge_point(fe4, "right"), color=RGBColor(0xA0, 0x30, 0x30))
arrow(edge_point(be6, "left"), edge_point(fe_yes, "right"), color=NAVY)

prs.save("activity_flow_slide.pptx")
print("saved")
