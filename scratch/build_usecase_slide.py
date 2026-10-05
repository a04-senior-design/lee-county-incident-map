from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn

NAVY   = RGBColor(0x1E, 0x3A, 0x5F)
GREY   = RGBColor(0x44, 0x44, 0x44)
BLUE_FILL  = RGBColor(0xDC, 0xE6, 0xF1)
GREEN_FILL = RGBColor(0xE2, 0xEF, 0xDA)
BLACK      = RGBColor(0x20, 0x20, 0x20)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
DARK_RED   = RGBColor(0xA0, 0x30, 0x30)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])


def add_shape(shape_type, left, top, width, height, text="", fill=WHITE,
              font_size=10, bold=False, font_color=BLACK, line_color=BLACK, line_width=1.25):
    shp = slide.shapes.add_shape(shape_type, Inches(left), Inches(top), Inches(width), Inches(height))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line_color
    shp.line.width = Pt(line_width)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(45720); tf.margin_right = Emu(45720)
    tf.margin_top = Emu(9000); tf.margin_bottom = Emu(9000)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(font_size)
        p.font.bold = bold
        p.font.color.rgb = font_color
    return shp


def edge_point(shp, side):
    l, t = shp.left / 914400, shp.top / 914400
    w, h = shp.width / 914400, shp.height / 914400
    if side == "top":    return (l + w / 2, t)
    if side == "bottom": return (l + w / 2, t + h)
    if side == "left":   return (l, t + h / 2)
    if side == "right":  return (l + w, t + h / 2)


def plain_line(p1, p2, color=BLACK, width=1.0):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                       Inches(p1[0]), Inches(p1[1]),
                                       Inches(p2[0]), Inches(p2[1]))
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    return conn


def dashed_arrow(p1, p2, color=BLACK, width=1.0):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                       Inches(p1[0]), Inches(p1[1]),
                                       Inches(p2[0]), Inches(p2[1]))
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    conn.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    ln = conn.line._get_or_add_ln()
    tail = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
    ln.append(tail)
    return conn


def label(text, x, y, size=8, color=GREY, bold=False, width=0.9):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(0.2))
    tf = box.text_frame
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    tf.text = text
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(size); p.font.italic = True; p.font.bold = bold; p.font.color.rgb = color


# ── Title ─────────────────────────────────────────────────────────────────
title_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.22), Inches(12.5), Inches(0.55))
tf = title_box.text_frame
tf.text = "Report Module — Use-Case Diagram"
p = tf.paragraphs[0]
p.font.size = Pt(25); p.font.bold = True; p.font.color.rgb = NAVY

sub_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.75), Inches(12.5), Inches(0.4))
tf2 = sub_box.text_frame
tf2.text = "Actors and «include» relationships, checked against what's actually implemented"
p2 = tf2.paragraphs[0]
p2.font.size = Pt(13); p2.font.italic = True; p2.font.color.rgb = GREY

# ── Left column: bullets ─────────────────────────────────────────────────
bullets_box = slide.shapes.add_textbox(Inches(0.35), Inches(1.30), Inches(3.85), Inches(5.9))
btf = bullets_box.text_frame
btf.word_wrap = True

sections = [
    ("Use-Case Structure", [
        "Registered User is the sole actor and directly initiates 5 use cases.",
        "Generate Report «includes» 5 backend-only operations executed inside "
        "one HTTP request — not user-visible steps.",
    ]),
    ("Gaps vs. Current Code", [
        "No auth enforced yet — app.py’s /api/report has a TODO to require a "
        "JWT; any visitor can generate a report today, registered or not.",
        "Type filter isn’t wired in — the nature/type checkboxes only affect "
        "the map view; POST /api/report only sends {days, bounds}, so use "
        "case 2 never reaches Generate Report.",
        "Date range is a single lookback slider (1–30 days), not yet the "
        "custom start/end range the design implies.",
        "Use cases 6–10 (bounds filter, date filter, map snapshot, charts, "
        "PDF assembly) match report.py / app.py exactly.",
    ]),
]

first = True
for heading, items in sections:
    hp = btf.paragraphs[0] if first else btf.add_paragraph()
    first = False
    hp.text = heading
    hp.font.size = Pt(14); hp.font.bold = True; hp.font.color.rgb = NAVY
    hp.space_after = Pt(4)
    for item in items:
        ip = btf.add_paragraph()
        ip.text = "•  " + item
        ip.font.size = Pt(10); ip.font.color.rgb = GREY
        ip.space_after = Pt(6)
    if heading != sections[-1][0]:
        sp = btf.add_paragraph()
        sp.space_after = Pt(4)

# ── Right column: use-case diagram ──────────────────────────────────────────
BOX_L, BOX_T, BOX_W, BOX_H = 5.3, 1.30, 7.55, 5.85
box = add_shape(MSO_SHAPE.RECTANGLE, BOX_L, BOX_T, BOX_W, BOX_H, "", WHITE, line_width=1.5)

box_title = slide.shapes.add_textbox(Inches(BOX_L), Inches(BOX_T + 0.06), Inches(BOX_W), Inches(0.35))
btf2 = box_title.text_frame
btf2.text = "Lee County Incident Map – Report Module"
bp = btf2.paragraphs[0]
bp.alignment = PP_ALIGN.CENTER
bp.font.size = Pt(14); bp.font.bold = True; bp.font.color.rgb = NAVY

# Actor (simple UML stick figure)
ax = 4.85  # actor center x
head = add_shape(MSO_SHAPE.OVAL, ax - 0.14, 3.55, 0.28, 0.28, "", WHITE, line_width=1.5)
plain_line((ax, 3.83), (ax, 4.35), width=1.5)                 # body
plain_line((ax - 0.35, 3.98), (ax + 0.35, 3.98), width=1.5)   # arms
plain_line((ax, 4.35), (ax - 0.27, 4.78), width=1.5)          # left leg
plain_line((ax, 4.35), (ax + 0.27, 4.78), width=1.5)          # right leg
label("Registered User", ax - 0.75, 4.82, size=10, color=BLACK, bold=False, width=1.5)

# Blue use cases (actor-initiated)
uc1 = add_shape(MSO_SHAPE.OVAL, 5.75, 1.78, 2.2, 0.62, "1. Configure\nDate Range", BLUE_FILL)
uc2 = add_shape(MSO_SHAPE.OVAL, 5.75, 2.58, 2.2, 0.62, "2. Configure Incident\nType Filter", BLUE_FILL, 9)
uc3 = add_shape(MSO_SHAPE.OVAL, 5.75, 3.38, 2.2, 0.55, "3. Set Viewport", BLUE_FILL)
uc4 = add_shape(MSO_SHAPE.OVAL, 5.75, 4.15, 2.2, 0.55, "4. Generate Report", BLUE_FILL, 10, True)
uc5 = add_shape(MSO_SHAPE.OVAL, 5.75, 4.92, 2.2, 0.55, "5. Download PDF", BLUE_FILL)

# Green use cases (included backend operations)
uc6  = add_shape(MSO_SHAPE.OVAL, 10.05, 1.78, 2.35, 0.62, "6. Filter Incidents\nby Bounds", GREEN_FILL, 9)
uc7  = add_shape(MSO_SHAPE.OVAL, 10.05, 2.58, 2.35, 0.62, "7. Filter Incidents\nby Date", GREEN_FILL, 9)
uc8  = add_shape(MSO_SHAPE.OVAL, 10.05, 3.38, 2.35, 0.62, "8. Render Map\nSnapshot", GREEN_FILL, 9)
uc9  = add_shape(MSO_SHAPE.OVAL, 10.05, 4.18, 2.35, 0.55, "9. Generate Charts", GREEN_FILL, 9)
uc10 = add_shape(MSO_SHAPE.OVAL, 10.05, 4.95, 2.35, 0.55, "10. Assemble PDF", GREEN_FILL, 9)

# Actor -- use case associations (plain lines)
actor_pt = (ax + 0.35, 3.98)
for uc in (uc1, uc2, uc3, uc4, uc5):
    plain_line(actor_pt, edge_point(uc, "left"))

# Generate Report --include--> backend operations (dashed arrows)
src = edge_point(uc4, "right")
for uc in (uc6, uc7, uc8, uc9, uc10):
    dst = edge_point(uc, "left")
    dashed_arrow(src, dst)
    mx, my = (src[0] + dst[0]) / 2, (src[1] + dst[1]) / 2
    label("«include»", mx - 0.45, my - 0.14, size=8)

prs.save("usecase_diagram_slide.pptx")
print("saved")
