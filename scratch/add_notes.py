from pptx import Presentation

NOTES = {
    "report_design_slide.pptx": (
        "This slide covers how the registered-user PDF report actually gets built, server-side. "
        "The whole thing runs in Flask, in memory, inside a single request — no separate service, "
        "no temp files, no external process.\n\n"
        "I picked ReportLab's Platypus layer over WeasyPrint because the report content is mixed — "
        "a map image, two chart images, a couple of tables — and Platypus lets me assemble that "
        "programmatically from BytesIO buffers, without an HTML template step or a system dependency "
        "like Pango.\n\n"
        "For the map snapshot, I first looked at browser-side capture — leaflet-image, html2canvas, "
        "dom-to-image. They all hit the same wall: OpenStreetMap's tile server doesn't send CORS "
        "headers, so the browser can't read the tile pixels to composite an image — the 'tainted "
        "canvas' problem. Puppeteer would sidestep that by screenshotting from a headless Chromium on "
        "the server, but that's a lot of operational weight for a lightweight Flask app. staticmap "
        "solves it cleanly: it downloads OSM tiles server-side over plain Python requests, which "
        "isn't subject to browser CORS at all.\n\n"
        "Charts are Matplotlib running its non-interactive Agg backend, so there's no display "
        "environment needed — figures render straight into memory. The two charts on the right "
        "aren't mockups — they're real output from the project's own _bar_chart() and _line_chart() "
        "functions, run against sample data."
    ),
    "activity_flow_slide.pptx": (
        "This is the request lifecycle for what happens when someone clicks 'Generate Report.'\n\n"
        "Design-wise, I kept it a single synchronous request/response — read the map bounds and date "
        "filter, POST once to /api/report, and everything downstream (filtering, rendering, PDF "
        "assembly) happens inside that one HTTP call. No polling, no websocket.\n\n"
        "Because it's synchronous, the button has to guard against a double-click while the report "
        "is generating — so the frontend disables it and relabels it 'Generating…' the instant it's "
        "clicked.\n\n"
        "On the backend, there's a fail-fast check: if the viewport and date filters together leave "
        "zero incidents, the endpoint returns HTTP 400 immediately, before doing any of the expensive "
        "work — downloading map tiles, rendering charts. No point building a report that would come "
        "back empty.\n\n"
        "Whichever way it goes, success or failure, both paths converge on the same 'restore button' "
        "step, wrapped in a finally block in the JavaScript, so the button can never get stuck loading. "
        "The diagram mirrors that directly: frontend on the left, backend on the right, and the two "
        "branches — bad request, and successful PDF — both flow back to that one recovery point on "
        "the frontend."
    ),
    "usecase_diagram_slide.pptx": (
        "This is the use-case diagram for the report subsystem. Registered User is the only actor, "
        "and they directly trigger five use cases: set the date range, filter by incident type, "
        "pan/zoom the viewport, generate the report, and download the PDF.\n\n"
        "Generate Report then 'includes' five backend operations — filter by bounds, filter by date, "
        "render the map snapshot, generate the charts, assemble the PDF. Those five aren't something "
        "the user watches happen one at a time; they're backend responsibilities that all fire inside "
        "that one request from the previous slide.\n\n"
        "The part I want to flag is the right column under 'Gaps vs. Current Code,' because this "
        "diagram describes the target design, not exactly where the code is today. Two real gaps: "
        "first, there's no auth check on the report endpoint yet — there's a literal TODO in app.py "
        "to require a JWT — so right now anyone can generate a report, not just a registered user. "
        "Second, the incident-type filter checked on the map doesn't actually get sent when you "
        "generate a report — the request body only carries the date window and bounding box, so a "
        "report today always includes every incident type regardless of what's toggled on the map.\n\n"
        "Everything else — the five included backend operations — is accurate and matches the code "
        "exactly."
    ),
}

for fname, text in NOTES.items():
    prs = Presentation(fname)
    slide = prs.slides[0]
    notes_tf = slide.notes_slide.notes_text_frame
    notes_tf.text = text
    prs.save(fname)
    print("updated", fname)
